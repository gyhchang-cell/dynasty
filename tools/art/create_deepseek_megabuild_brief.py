from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.7)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Songti SC'
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Songti SC')
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.15

pages = [
('DeepSeek 王朝大型建筑开发任务', [
('工作目标', '请直接在 /Users/a15356015027/Desktop/dynasty 开发。先完整阅读 docs/开发约定.md 和现有结构实现。Codex 同时修复武器模型、龙长度、装备图标、方块贴图和进化任务。你负责大型建筑与建筑内部探索，交付可以运行的代码及验证结果，不只写方案。'),
('第一项 巨型天工山城', '新增主世界结构 dynasty:tiangong_citadel。建议占地 176×176 格、高约 72 格，至少 100000 个最终非空气方块。统计不能把空气、重复覆盖或纯实心地下填充计入主体指标。城墙、层叠殿阁、矿场、仓储和街巷组成完整建筑群，不能只是把小房子同比例放大。'),
('外观和动线', '以中国王朝城寨为主题，原版深板岩、石砖、木材、铜、灯笼为主。主轴由山门、内街、工坊院、主殿组成；左右院落错落，至少三个可达高度层。屋顶用楼梯和台阶形成出檐，加入梁柱、窗棂、女儿墙、回廊、塔楼和桥梁。入口、楼梯与每层出口连续可走，普通玩家不飞行、不拆墙能走完整条路线。'),
('可以玩的内容', '至少十二个有用途的空间：采石场、矿坑、冶炼房、兵器坊、粮仓、药圃、藏书楼、守卫营、密库、主殿及两处支路探索点。主线路程约十至十五分钟，支路能获得真实的现有淬炼材料、锻造材料和有用战利品。用实际注册物品，不造不存在的奖励 ID；高阶材料数量克制，不让早期直接毕业。'),
('执行边界', '新增代码统一放 com.dynasty.structure.megabuild 包；新增数据统一使用 tiangong_ 前缀；工具放 tools/megabuild，文档放 docs/megabuild。公共注册文件 DynastyStructures.java 只在交付目录写出最小接入补丁，不直接改它，交给 Codex 整合。不要改已有建筑、玩家存档、贴图、语言总表、武器、任务生成器和部署脚本。'),
]),
('世界生成和第二座资源建筑', [
('跨区块可靠性', '必须按区块裁剪放置，不能在每个区块重复遍历整座十万方块建筑。大建筑拆成合理部件，所有局部坐标落在声明包围盒内，覆盖区块边界和旋转测试。不要在生成回调里强制加载其他区块。序列化保存所有布局参数，重载以后形状不变。'),
('落地和性能', '平原或山地缓坡生成，避开海底和高度越界；多点采样地形并设高差上限。允许分层台基，禁止生成直插世界底部的巨大实心柱。结构自然生成要稀疏，建议 spacing 96、separation 48 并检查单位与当前实现。怪物使用已有实体，数量有界，不做无限刷怪器，不让整个城同时激活几十个 Boss。'),
('第二项 山麓采矿庄园', '主城完成并验证后再新增 tiangong_mining_estate，建议 96×96 格，含露天采石场、三层矿坑、冶炼工坊、住宅和材料仓库。它提供中期材料收集玩法，与主城不同，不复制主城缩小版。原版矿石数量合理，放在可探索矿脉和宝箱，不铺满钻石。新增结构类型与部件仍放独立包，接入补丁一起交付。'),
('战利品与发现入口', '给不同房间设置不同战利品表，普通储物间以食物、火把、铁铜、低阶素材为主；深处密库才出现稀有材料。无刷新复制漏洞。准备结构中文名、英文名、探险者指南针显示词条和进入结构的 advancement 文件；语言键单独交付 JSON，不能改总表。'),
('探索任务素材', '交付 docs/megabuild/quest-proposals.json，给两座建筑各写六个有实质内容的节点，包括入口、收集材料、探索支路和最终宝库。每个节点列出现有 item ID、数量、建议前置、完成检测方式及具体收益。只能使用现有真实检测，不把勾选说明伪装成自动探索检测。由 Codex 接入现有任务树。'),
]),
('验证和交付', [
('必须通过的检查', '单独编译新增类并给出复现命令；不能为了消除报错修改共享文件或降低已有测试。验证包围盒、四种朝向、跨区块放置、NBT 保存加载、楼梯上下连续、至少两格净空、宝箱不堵路。计数统计最终非空气方块，另列台基填充、建筑主体、装饰与空气体积。'),
('视觉检查', '输出真实方块数据生成的俯视图、正立面和斜视图，标明这是否离线预览。不要拿概念图冒充游戏效果。屋顶、门窗、桥和塔楼必须有清楚层次；主城至少十万实体方块这个目标不能用大量看不见的地下填充凑数。尽可能输出可导入测试世界的结构或明确放置命令，但不触碰当前玩家的存档。'),
('交付文件', '在 docs/megabuild/handoff.md 写明：新增文件清单、接入补丁位置、结构 ID、尺寸、实方块统计、奖励表、复现步骤、编译结果及未验证项。公共注册补丁和语言增量分别保存，Codex 合并时能精确定位。修改前记录现有工作区状态，不覆盖同伴更新，不执行 git reset、不批量恢复文件、不提交他人的修改。'),
('工作顺序', '先完成天工山城的外观与通路，再加入房间和奖励，验证跨区块生成，最后做采矿庄园。每完成一个阶段就更新 handoff.md。遇到接入点不明确时先读项目已有 Structure、Piece、Loot 实现，按现有版本 API 开发；不要凭记忆写不存在的方法。尽量持续完成上述范围，不扩展为新战斗系统。'),
('与 Codex 的分工', '你完成这份任务后，把 handoff.md 的路径发给用户。Codex 负责公共注册、语言合并、指南针与任务接入、最终编译打包。当前没有要求你重绘美术，也没有要求你动正在运行的游戏。请从检查现有结构代码开始直接实施。'),
])]
for n, (title, sections) in enumerate(pages):
    if n: doc.add_page_break()
    doc.add_paragraph(title, 'Title' if n == 0 else 'Heading 1')
    for heading, body in sections:
        doc.add_paragraph(heading, 'Heading 2')
        doc.add_paragraph(body)
for p in doc.paragraphs:
    pp = p._p.get_or_add_pPr()
    for border in list(pp.findall(qn('w:pBdr'))): pp.remove(border)
for style in doc.styles:
    pp = style.element.find(qn('w:pPr'))
    if pp is not None:
        for border in list(pp.findall(qn('w:pBdr'))): pp.remove(border)
doc.save('/Users/a15356015027/Desktop/DeepSeek王朝大型建筑任务.docx')
